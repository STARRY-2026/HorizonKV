"""Create the editable PowerPoint teaser figure for MemHeadKV."""

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "results" / "cvpr" / "figure1"
PPT_PATH = OUT_DIR / "figure1_memheadkv.pptx"
PREVIEW_PATH = OUT_DIR / "figure1_memheadkv_preview.png"
SCORES_PATH = (
    ROOT
    / "results"
    / "myobservations"
    / "effective_memory_readout_core_top100_n50_o5000"
    / "effective_readout_scores.csv"
)
FRAME_DIR = OUT_DIR / "assets"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = "17324D"
BLUE = "3568A8"
TEAL = "168C86"
TEAL_DARK = "0B6864"
TEAL_LIGHT = "DDF3F0"
CORAL = "DE6655"
CORAL_LIGHT = "FBE8E3"
GOLD = "E5A832"
GOLD_LIGHT = "FFF1CE"
INK = "24323F"
MUTED = "647481"
GRID = "DCE4E8"
PANEL = "F7F9FA"
WHITE = "FFFFFF"
LIGHT = "EFF3F5"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=10,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    font="Arial",
    margin=0.02,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x, y, w, h, fill, line=GRID, radius=False, line_width=0.8):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_circle(slide, x, y, diameter, fill, line=WHITE, line_width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_line(slide, x1, y1, x2, y2, color=MUTED, width=1.2, dash=None):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_arrow(slide, x, y, w, h, fill=BLUE):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = rgb(fill)
    arrow.line.fill.background()
    return arrow


def add_picture_crop(slide, image_path, x, y, w, h):
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    target_ratio = w / h
    picture = slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > target_ratio:
        visible = target_ratio / image_ratio
        picture.crop_left = picture.crop_right = (1.0 - visible) / 2.0
    else:
        visible = image_ratio / target_ratio
        picture.crop_top = picture.crop_bottom = (1.0 - visible) / 2.0
    return picture


def mix_color(low, high, fraction):
    low_rgb = np.asarray([int(low[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    high_rgb = np.asarray([int(high[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    value = np.rint(low_rgb + np.clip(fraction, 0, 1) * (high_rgb - low_rgb)).astype(int)
    return "".join(f"{component:02X}" for component in value)


def draw_stacked_bar(slide, label, y, memory_value, current_value):
    x = 4.63
    width = 3.23
    add_text(slide, label, x, y - 0.28, 1.3, 0.24, size=8.5, bold=True)
    memory_w = width * memory_value / 100.0
    add_rect(slide, x, y, memory_w, 0.30, TEAL, line=TEAL, radius=True)
    add_rect(slide, x + memory_w, y, width - memory_w, 0.30, GOLD_LIGHT, line=GOLD_LIGHT, radius=True)
    add_text(
        slide,
        f"Memory {memory_value:.1f}%",
        x + 0.05,
        y,
        max(memory_w - 0.08, 0.55),
        0.30,
        size=7.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        f"Current {current_value:.1f}%",
        x + memory_w + 0.03,
        y,
        max(width - memory_w - 0.06, 0.58),
        0.30,
        size=7.2,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def extract_frames():
    from decord import VideoReader

    FRAME_DIR.mkdir(parents=True, exist_ok=True)
    video_path = "/data/streamingbench/videos/sample_17_real.mp4"
    reader = VideoReader(video_path, num_threads=1)
    fps = float(reader.get_avg_fps())
    times = [300, 330, 350, 362]
    paths = []
    for index, timestamp in enumerate(times):
        frame_idx = min(len(reader) - 1, int(timestamp * fps))
        image = Image.fromarray(reader[frame_idx].asnumpy())
        path = FRAME_DIR / f"sample17_t{timestamp}.png"
        image.save(path)
        paths.append(path)
    return paths, times


def add_panel_header(slide, label, title, x, width, accent):
    add_circle(slide, x, 0.78, 0.28, accent, line=accent)
    add_text(
        slide,
        label,
        x,
        0.78,
        0.28,
        0.28,
        size=9,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, title, x + 0.38, 0.74, width - 0.38, 0.36, size=12.5, bold=True)


def draw_panel_a(slide, frame_paths, times):
    x0 = 0.34
    add_panel_header(slide, "a", "A question arrives after the stream", x0, 4.05, BLUE)
    thumb_y = 1.30
    thumb_w = 0.87
    thumb_h = 0.82
    gap = 0.10
    for index, (path, timestamp) in enumerate(zip(frame_paths, times)):
        x = x0 + index * (thumb_w + gap)
        add_picture_crop(slide, path, x, thumb_y, thumb_w, thumb_h)
        border = add_rect(slide, x, thumb_y, thumb_w, thumb_h, WHITE, line=TEAL if index == 0 else GRID)
        border.fill.transparency = 100
        border.line.width = Pt(2.2 if index == 0 else 0.8)
        add_text(
            slide,
            f"t={timestamp}s",
            x,
            thumb_y + thumb_h + 0.02,
            thumb_w,
            0.18,
            size=6.7,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    add_arrow(slide, x0, 2.31, 3.78, 0.16, fill=LIGHT)
    add_text(slide, "streaming video history", x0 + 0.84, 2.26, 2.0, 0.20, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)

    add_rect(slide, x0 + 0.26, 2.69, 3.50, 0.66, WHITE, line=BLUE, radius=True, line_width=1.2)
    add_text(
        slide,
        'Question at t_q: "What is the man using\nto cook food in the pan?"',
        x0 + 0.43,
        2.75,
        3.16,
        0.52,
        size=9,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_text(slide, "Boundary-responsive Head", x0 + 0.05, 3.64, 1.62, 0.25, size=8.5, bold=True, color=CORAL)
    add_text(slide, "Historical Evidence Head", x0 + 0.05, 4.60, 1.62, 0.25, size=8.5, bold=True, color=TEAL_DARK)
    for row_y, color in [(4.02, CORAL), (4.98, TEAL)]:
        for index in range(5):
            add_circle(slide, x0 + 1.72 + index * 0.41, row_y, 0.20, LIGHT, line=GRID, line_width=0.5)
        add_circle(slide, x0 + 1.72, row_y, 0.20, color, line=color)
    add_line(slide, x0 + 1.82, 5.08, x0 + 0.44, 1.73, color=TEAL, width=2.2)
    add_circle(slide, x0 + 0.36, 1.65, 0.16, TEAL, line=WHITE, line_width=1.4)
    add_line(slide, x0 + 3.46, 4.12, x0 + 3.31, 2.13, color=CORAL, width=2.0)
    add_circle(slide, x0 + 3.23, 2.04, 0.16, CORAL, line=WHITE, line_width=1.4)
    add_text(slide, "clean internal evidence", x0 + 0.13, 5.29, 1.63, 0.22, size=7.2, color=TEAL_DARK, bold=True)
    add_text(slide, "chunk boundary / recent", x0 + 2.22, 5.29, 1.66, 0.22, size=7.2, color=CORAL, bold=True, align=PP_ALIGN.RIGHT)
    add_rect(slide, x0 + 0.25, 5.77, 3.52, 0.54, TEAL_LIGHT, line=TEAL_LIGHT, radius=True)
    add_text(
        slide,
        "Seeing is not remembering.",
        x0 + 0.37,
        5.82,
        3.28,
        0.42,
        size=12,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def draw_panel_b(slide, scores):
    x0 = 4.48
    add_panel_header(slide, "b", "Historical readout is concentrated", x0, 4.18, TEAL)
    add_text(slide, "Eager pseudo-query probe", x0 + 0.15, 1.20, 1.8, 0.25, size=8.5, color=MUTED, bold=True)
    add_text(slide, "203 streaming observations", x0 + 1.80, 1.20, 1.95, 0.25, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)
    draw_stacked_bar(slide, "Visual tokens", 1.67, 25.9, 74.1)
    draw_stacked_bar(slide, "Probe readout", 2.26, 75.8, 24.2)
    add_text(
        slide,
        "A compact retained memory receives disproportionate readout.",
        x0 + 0.15,
        2.74,
        3.75,
        0.30,
        size=8,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_text(slide, "True-question HE score", x0 + 0.15, 3.19, 1.85, 0.25, size=8.5, color=MUTED, bold=True)
    add_text(
        slide,
        "112 physical KV Heads = 28 layers x 4 Heads",
        x0 + 1.78,
        3.19,
        2.16,
        0.25,
        size=7.2,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    matrix = np.full((28, 4), np.nan)
    for row in scores.itertuples():
        matrix[int(row.layer), int(row.kv_head)] = float(row.internal_top100_mean_attention)
    low, high = np.nanquantile(matrix, [0.05, 0.95])
    grid_x = x0 + 0.50
    grid_y = 3.57
    cell_w = 0.47
    cell_h = 0.083
    top_indices = set(map(tuple, np.argwhere(matrix >= np.nanquantile(matrix, 0.90))))
    for layer in range(28):
        for head in range(4):
            value = matrix[layer, head]
            fraction = (value - low) / max(high - low, 1e-12)
            fill = mix_color("FFF2D3", "C93C51", fraction)
            cell = add_rect(
                slide,
                grid_x + head * cell_w,
                grid_y + layer * cell_h,
                cell_w,
                cell_h,
                fill,
                line=WHITE,
                line_width=0.25,
            )
            if (layer, head) in top_indices:
                cell.line.color.rgb = rgb(TEAL_DARK)
                cell.line.width = Pt(0.8)
    for head in range(4):
        add_text(
            slide,
            f"K{head}",
            grid_x + head * cell_w,
            grid_y - 0.22,
            cell_w,
            0.18,
            size=6.5,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    add_text(slide, "L0", grid_x - 0.31, grid_y - 0.01, 0.27, 0.18, size=6.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "L27", grid_x - 0.34, grid_y + 27 * cell_h - 0.02, 0.30, 0.18, size=6.5, color=MUTED, align=PP_ALIGN.RIGHT)

    add_rect(slide, x0 + 2.72, 3.57, 1.22, 0.42, CORAL_LIGHT, line=CORAL_LIGHT, radius=True)
    add_text(slide, "high HE response", x0 + 2.78, 3.62, 1.10, 0.30, size=7.5, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, x0 + 2.72, 4.14, 1.22, 0.42, LIGHT, line=LIGHT, radius=True)
    add_text(slide, "weak response", x0 + 2.78, 4.19, 1.10, 0.30, size=7.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "M_g = E_q [ Top100Mean ( A_g^q ( H_internal ) ) ]",
        x0 + 2.35,
        4.77,
        1.75,
        0.64,
        size=8.5,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, x0 + 2.39, 5.53, 1.66, 0.74, WHITE, line=GRID, radius=True)
    add_text(slide, "Remove boundary windows", x0 + 2.48, 5.58, 1.48, 0.22, size=7.3, bold=True, align=PP_ALIGN.CENTER)
    segments = [(0.0, 0.25, CORAL_LIGHT), (0.25, 0.83, TEAL_LIGHT), (0.83, 1.0, CORAL_LIGHT)]
    for start, end, fill in segments:
        add_rect(slide, x0 + 2.61 + start * 1.22, 5.91, (end - start) * 1.22, 0.14, fill, line=fill)
    add_text(slide, "boundary", x0 + 2.44, 6.08, 0.58, 0.18, size=5.8, color=CORAL, align=PP_ALIGN.CENTER)
    add_text(slide, "internal history", x0 + 2.95, 6.08, 0.72, 0.18, size=5.8, color=TEAL_DARK, align=PP_ALIGN.CENTER)


def draw_cache_rows(slide, x, y, widths, fill, label_prefix):
    max_width = 1.20
    for index, width_fraction in enumerate(widths):
        label = label_prefix[:-1] if label_prefix.endswith(":") else f"{label_prefix}{index + 1}"
        add_text(slide, label, x, y + index * 0.39, 0.32, 0.25, size=6.8, color=MUTED, align=PP_ALIGN.RIGHT)
        add_rect(slide, x + 0.38, y + index * 0.39, max_width * width_fraction, 0.23, fill, line=fill, radius=True)
        for tick in range(max(1, int(7 * width_fraction))):
            tx = x + 0.43 + tick * (max_width * width_fraction - 0.10) / max(1, int(7 * width_fraction) - 1)
            add_circle(slide, tx, y + index * 0.39 + 0.065, 0.07, WHITE, line=WHITE, line_width=0.3)


def draw_panel_c(slide):
    x0 = 8.67
    add_panel_header(slide, "c", "MemHeadKV: evidence-aware allocation", x0, 4.26, CORAL)
    add_text(slide, "Uniform HERMES", x0 + 0.15, 1.33, 1.55, 0.28, size=9, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    draw_cache_rows(slide, x0 + 0.18, 1.75, [0.72] * 6, BLUE, "H")
    add_arrow(slide, x0 + 1.86, 2.54, 0.54, 0.24, fill=LIGHT)
    add_text(slide, "HE score", x0 + 1.82, 2.21, 0.62, 0.24, size=7, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Same total B", x0 + 1.72, 2.86, 0.84, 0.22, size=6.7, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "MemHeadKV", x0 + 2.51, 1.33, 1.38, 0.28, size=9, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
    heterogeneous = [1.00, 0.50, 0.84, 0.42, 1.00, 0.56]
    colors = [TEAL, GOLD, TEAL, GOLD, TEAL, GOLD]
    for index, (width, fill) in enumerate(zip(heterogeneous, colors)):
        draw_cache_rows(slide, x0 + 2.50, 1.75 + index * 0.39, [width], fill, f"H{index + 1}:")

    add_rect(slide, x0 + 0.19, 4.28, 3.72, 0.68, PANEL, line=GRID, radius=True)
    add_text(slide, "Offline profile", x0 + 0.30, 4.38, 0.84, 0.20, size=7.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, x0 + 1.16, 4.43, 0.38, 0.13, fill=GRID)
    add_text(slide, "Head-wise budget", x0 + 1.58, 4.38, 1.05, 0.20, size=7.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, x0 + 2.68, 4.43, 0.38, 0.13, fill=GRID)
    add_text(slide, "Ragged KV cache", x0 + 3.07, 4.38, 0.75, 0.20, size=7.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    badges = [
        ("Training-free", TEAL_LIGHT, TEAL_DARK),
        ("Same budget", GOLD_LIGHT, "8A6417"),
        ("Physical KV Heads", CORAL_LIGHT, CORAL),
    ]
    badge_x = x0 + 0.24
    for text, fill, color in badges:
        width = 1.08 if text != "Physical KV Heads" else 1.34
        add_rect(slide, badge_x, 5.27, width, 0.38, fill, line=fill, radius=True)
        add_text(slide, text, badge_x + 0.03, 5.31, width - 0.06, 0.28, size=7.2, color=color, bold=True, align=PP_ALIGN.CENTER)
        badge_x += width + 0.12

    add_rect(slide, x0 + 0.28, 5.91, 3.58, 0.42, WHITE, line=CORAL, radius=True, line_width=1.1)
    add_text(
        slide,
        "Protect the Heads that retrieve historical evidence.",
        x0 + 0.38,
        5.96,
        3.38,
        0.31,
        size=9.2,
        color=CORAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def make_preview(frame_paths, scores):
    width, height = 2400, 1350
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font_title = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 46)
        font_header = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 31)
        font_body = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 22)
        font_small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 18)
    except OSError:
        font_title = font_header = font_body = font_small = ImageFont.load_default()

    draw.text((62, 35), "Not All Visual Heads Remember", font=font_title, fill="#17324D")
    draw.text((62, 92), "Historical Evidence Heads for Streaming Video Understanding", font=font_body, fill="#647481")
    panels = [(60, 160, 770, 1175), (815, 160, 1555, 1175), (1600, 160, 2340, 1175)]
    headers = [
        ("a", "A question arrives after the stream", "#3568A8"),
        ("b", "Historical readout is concentrated", "#168C86"),
        ("c", "MemHeadKV: evidence-aware allocation", "#DE6655"),
    ]
    for panel, (label, title, accent) in zip(panels, headers):
        x1, y1, x2, y2 = panel
        draw.rounded_rectangle(panel, radius=12, fill="white", outline="white", width=1)
        draw.ellipse((x1 + 18, y1 + 20, x1 + 64, y1 + 66), fill=accent)
        draw.text((x1 + 34, y1 + 25), label, font=font_body, fill="white", anchor="ma")
        draw.text((x1 + 82, y1 + 22), title, font=font_header, fill="#24323F")

    x1, y1, _, _ = panels[0]
    for index, path in enumerate(frame_paths):
        frame = Image.open(path).convert("RGB")
        frame.thumbnail((155, 116))
        image.paste(frame, (x1 + 28 + index * 175, y1 + 120))
    draw.rounded_rectangle((x1 + 80, y1 + 295, x1 + 690, y1 + 405), radius=18, fill="white", outline="#3568A8", width=3)
    draw.multiline_text((x1 + 385, y1 + 350), 'Question: "What is the man using\nto cook food in the pan?"', font=font_body, fill="#17324D", anchor="mm", align="center")
    draw.text((x1 + 385, y1 + 530), "Boundary-responsive Head", font=font_body, fill="#DE6655", anchor="mm")
    draw.text((x1 + 385, y1 + 665), "Historical Evidence Head", font=font_body, fill="#0B6864", anchor="mm")
    draw.line((x1 + 600, y1 + 550, x1 + 600, y1 + 225), fill="#DE6655", width=8)
    draw.line((x1 + 250, y1 + 685, x1 + 110, y1 + 225), fill="#168C86", width=8)
    draw.rounded_rectangle((x1 + 95, y1 + 855, x1 + 675, y1 + 945), radius=18, fill="#DDF3F0")
    draw.text((x1 + 385, y1 + 900), "Seeing is not remembering.", font=font_header, fill="#0B6864", anchor="mm")

    x1, y1, _, _ = panels[1]
    draw.text((x1 + 35, y1 + 115), "Visual tokens", font=font_body, fill="#24323F")
    draw.rectangle((x1 + 35, y1 + 155, x1 + 210, y1 + 205), fill="#168C86")
    draw.rectangle((x1 + 210, y1 + 155, x1 + 680, y1 + 205), fill="#FFF1CE")
    draw.text((x1 + 35, y1 + 235), "Probe readout", font=font_body, fill="#24323F")
    draw.rectangle((x1 + 35, y1 + 275, x1 + 545, y1 + 325), fill="#168C86")
    draw.rectangle((x1 + 545, y1 + 275, x1 + 680, y1 + 325), fill="#FFF1CE")
    draw.text((x1 + 120, y1 + 180), "25.9%", font=font_small, fill="white", anchor="mm")
    draw.text((x1 + 290, y1 + 300), "75.8% memory readout", font=font_small, fill="white", anchor="mm")
    matrix = np.full((28, 4), np.nan)
    for row in scores.itertuples():
        matrix[int(row.layer), int(row.kv_head)] = float(row.internal_top100_mean_attention)
    low, high = np.nanquantile(matrix, [0.05, 0.95])
    gx, gy, cw, ch = x1 + 75, y1 + 405, 70, 18
    for layer in range(28):
        for head in range(4):
            value = (matrix[layer, head] - low) / max(high - low, 1e-12)
            fill = "#" + mix_color("FFF2D3", "C93C51", value)
            draw.rectangle((gx + head * cw, gy + layer * ch, gx + (head + 1) * cw - 1, gy + (layer + 1) * ch - 1), fill=fill)
    draw.multiline_text((x1 + 500, y1 + 620), "True-question HE score\n28 layers x 4 KV Heads\n\nBoundary-cleaned\ninternal history", font=font_body, fill="#17324D", spacing=12, align="center", anchor="mm")

    x1, y1, _, _ = panels[2]
    draw.text((x1 + 165, y1 + 125), "Uniform HERMES", font=font_body, fill="#17324D", anchor="mm")
    draw.text((x1 + 555, y1 + 125), "MemHeadKV", font=font_body, fill="#DE6655", anchor="mm")
    for index in range(6):
        yy = y1 + 185 + index * 65
        draw.rounded_rectangle((x1 + 45, yy, x1 + 285, yy + 35), radius=12, fill="#3568A8")
    widths = [300, 150, 252, 126, 300, 168]
    colors = ["#168C86", "#E5A832"] * 3
    for index, (bar_width, color) in enumerate(zip(widths, colors)):
        yy = y1 + 185 + index * 65
        draw.rounded_rectangle((x1 + 400, yy, x1 + 400 + bar_width, yy + 35), radius=12, fill=color)
    draw.text((x1 + 365, y1 + 395), "same B", font=font_small, fill="#0B6864", anchor="mm")
    draw.rounded_rectangle((x1 + 60, y1 + 690, x1 + 680, y1 + 800), radius=18, fill="#EFF3F5")
    draw.text((x1 + 370, y1 + 745), "Offline profile  >  Head-wise budget  >  Ragged KV", font=font_body, fill="#17324D", anchor="mm")
    draw.rounded_rectangle((x1 + 70, y1 + 870, x1 + 670, y1 + 950), radius=18, fill="white", outline="#DE6655", width=3)
    draw.text((x1 + 370, y1 + 910), "Protect Heads that retrieve historical evidence.", font=font_body, fill="#DE6655", anchor="mm")

    draw.text((60, 1240), "Figure 1. Discovery and system overview of Historical Evidence Heads and MemHeadKV.", font=font_body, fill="#24323F")
    image.save(PREVIEW_PATH)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    frame_paths, times = extract_frames()
    scores = pd.read_csv(SCORES_PATH)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(WHITE)

    add_text(slide, "Not All Visual Heads Remember", 0.34, 0.14, 5.25, 0.39, size=20, color=NAVY, bold=True)
    add_text(
        slide,
        "Discovering Historical Evidence Heads for Streaming Video Understanding",
        5.60,
        0.18,
        7.37,
        0.32,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    add_line(slide, 0.34, 0.62, 12.99, 0.62, color=GRID, width=1.0)

    for x, width in [(0.24, 4.04), (4.38, 4.08), (8.56, 4.53)]:
        add_rect(slide, x, 0.70, width, 5.83, WHITE, line=WHITE, radius=False, line_width=0.1)
    add_line(slide, 4.33, 0.83, 4.33, 6.38, color=GRID, width=0.9)
    add_line(slide, 8.51, 0.83, 8.51, 6.38, color=GRID, width=0.9)

    draw_panel_a(slide, frame_paths, times)
    draw_panel_b(slide, scores)
    draw_panel_c(slide)

    add_text(
        slide,
        "Figure 1. Historical Evidence Heads selectively read boundary-cleaned internal video memory under real questions; MemHeadKV converts this stable head heterogeneity into unequal physical KV-cache budgets without training or increasing the average cache budget.",
        0.40,
        6.70,
        12.53,
        0.54,
        size=8.5,
        color=INK,
        valign=MSO_ANCHOR.TOP,
    )

    presentation.save(PPT_PATH)
    make_preview(frame_paths, scores)
    print(PPT_PATH)
    print(PREVIEW_PATH)


if __name__ == "__main__":
    main()
