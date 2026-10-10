"""Build the compact, editable MemHeadKV Introduction figure as PowerPoint."""

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "results" / "cvpr" / "figure1"
ASSET_DIR = OUT_DIR / "intro_assets"
OUT_PATH = OUT_DIR / "figure1_memheadkv_intro_compact.pptx"
SCORES_PATH = (
    ROOT
    / "results/myobservations/effective_memory_readout_core_top100_n50_o5000"
    / "effective_readout_scores.csv"
)
VIDEO_PATH = "/data/streamingbench/videos/sample_17_real.mp4"

NAVY = "19334C"
INK = "2B3943"
MUTED = "71808A"
BLUE = "4A78B4"
TEAL = "14918A"
TEAL_DARK = "086B65"
TEAL_PALE = "DEF3F0"
ORANGE = "E98247"
ORANGE_DARK = "CB6431"
ORANGE_PALE = "FBE7D8"
CORAL = "D95F54"
PANEL = "F3F8FA"
LINE = "BDD0D7"
LIGHT = "F5F7F8"
WHITE = "FFFFFF"


def color(value):
    return RGBColor.from_string(value)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=9,
    fill=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.02,
    font="Arial",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return box


def add_rect(slide, x, y, w, h, fill, edge=LINE, rounded=False, width=0.7, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if transparency:
        shape.fill.transparency = transparency
    shape.line.color.rgb = color(edge)
    shape.line.width = Pt(width)
    return shape


def add_circle(slide, x, y, diameter, fill, edge=None, width=0.8):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(edge or fill)
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1, y1, x2, y2, fill=MUTED, width=1.0):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color(fill)
    line.line.width = Pt(width)
    return line


def add_arrow(slide, x, y, w, h, fill=BLUE):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = color(fill)
    arrow.line.fill.background()
    return arrow


def add_picture_crop(slide, path, x, y, w, h):
    with Image.open(path) as image:
        ratio = image.width / image.height
    target = w / h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if ratio > target:
        visible = target / ratio
        picture.crop_left = picture.crop_right = (1.0 - visible) / 2.0
    else:
        visible = ratio / target
        picture.crop_top = picture.crop_bottom = (1.0 - visible) / 2.0
    return picture


def extract_frames():
    from decord import VideoReader

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    times = [278, 286, 294, 302, 310, 318, 326, 334, 342, 350, 356, 362]
    reader = VideoReader(VIDEO_PATH, num_threads=1)
    fps = float(reader.get_avg_fps())
    paths = []
    for timestamp in times:
        index = min(len(reader) - 1, int(timestamp * fps))
        path = ASSET_DIR / f"sample17_{timestamp}s.png"
        Image.fromarray(reader[index].asnumpy()).save(path)
        paths.append(path)
    return paths, times


def gaussian(x, center, width, height):
    return height * np.exp(-0.5 * ((x - center) / width) ** 2)


def conceptual_curves(count, evidence_center):
    x = np.linspace(0, 1, count)
    he = 0.07 + 0.025 * (np.sin(15 * np.pi * x) + 1)
    other = 0.07 + 0.022 * (np.cos(13 * np.pi * x) + 1)
    he_specs = [
        (0.04, 0.025, 0.25), (0.11, 0.028, 0.42), (0.18, 0.023, 0.28),
        (evidence_center, 0.032, 0.90), (0.37, 0.024, 0.34), (0.45, 0.030, 0.47),
        (0.55, 0.023, 0.25), (0.64, 0.027, 0.38), (0.73, 0.023, 0.29),
        (0.82, 0.030, 0.45), (0.93, 0.025, 0.27),
    ]
    other_specs = [
        (0.03, 0.022, 0.34), (0.10, 0.032, 0.51), (0.19, 0.026, 0.28),
        (0.35, 0.028, 0.22), (0.44, 0.022, 0.31), (0.57, 0.027, 0.27),
        (0.68, 0.030, 0.37), (0.79, 0.033, 0.55), (0.89, 0.030, 0.82),
        (0.96, 0.022, 0.49),
    ]
    for center, width, height in he_specs:
        he += gaussian(x, center, width, height)
    for center, width, height in other_specs:
        other += gaussian(x, center, width, height)
    maximum = max(float(he.max()), float(other.max()))
    return he / maximum, other / maximum


def draw_editable_track(slide, x, y, w, h, values, stroke, pale):
    add_rect(slide, x, y, w, h, WHITE, edge="D8E4E8", rounded=True, width=0.65)
    column = w / len(values)
    points = []
    for index, value in enumerate(values):
        bar_h = max(0.025, (h - 0.12) * float(value))
        left = x + index * column
        add_rect(
            slide,
            left,
            y + h - 0.06 - bar_h,
            column + 0.004,
            bar_h,
            pale,
            edge=pale,
            width=0.1,
        )
        points.append((left + column / 2, y + h - 0.06 - bar_h))
    for first, second in zip(points, points[1:]):
        add_line(slide, first[0], first[1], second[0], second[1], fill=stroke, width=1.45)


def dashed_vertical(slide, x, y1, y2, fill=TEAL):
    dash_h = 0.07
    gap = 0.055
    y = y1
    while y < y2:
        add_rect(slide, x - 0.008, y, 0.016, min(dash_h, y2 - y), fill, edge=fill, width=0.1)
        y += dash_h + gap


def mix(low, high, fraction):
    a = np.asarray([int(low[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    b = np.asarray([int(high[i : i + 2], 16) for i in (0, 2, 4)], dtype=float)
    result = np.rint(a + np.clip(fraction, 0, 1) * (b - a)).astype(int)
    return "".join(f"{component:02X}" for component in result)


def panel_title(slide, label, title, x, y, accent):
    add_circle(slide, x, y, 0.27, accent)
    add_text(slide, label, x, y, 0.27, 0.27, size=8.5, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.35, y - 0.01, 3.55, 0.30, size=11.2, fill=NAVY, bold=True)


def draw_panel_a(slide, x0, frame_paths, times):
    panel_title(slide, "a", "Head-wise historical readout", x0 + 0.10, 0.82, BLUE)
    add_rect(slide, x0 + 0.10, 1.19, 3.82, 1.00, WHITE, edge=BLUE, rounded=True, width=1.0)
    add_circle(slide, x0 + 0.20, 1.30, 0.28, BLUE)
    add_text(slide, "Q", x0 + 0.20, 1.30, 0.28, 0.28, size=9.5, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "What is the man using to cook food in the pan?",
        x0 + 0.54,
        1.27,
        2.80,
        0.35,
        size=8.4,
        fill=INK,
        bold=True,
    )
    add_text(slide, "tq=363s", x0 + 3.28, 1.29, 0.52, 0.25, size=7.0, fill=BLUE, bold=True, align=PP_ALIGN.RIGHT)
    options = [
        ("A. wooden spoon", x0 + 0.55, 1.67, 1.33, LIGHT, "E0E7EA", MUTED),
        ("B. metal spatula", x0 + 1.96, 1.67, 1.33, LIGHT, "E0E7EA", MUTED),
        ("C. red kitchen tongs", x0 + 0.55, 1.94, 1.55, TEAL_PALE, TEAL, TEAL_DARK),
        ("D. fork", x0 + 2.18, 1.94, 1.11, LIGHT, "E0E7EA", MUTED),
    ]
    for label, x, y, w, fill, edge, text_fill in options:
        add_rect(slide, x, y, w, 0.21, fill, edge=edge, rounded=True, width=0.55)
        add_text(slide, label, x + 0.03, y, w - 0.06, 0.21, size=6.5, fill=text_fill, bold=label.startswith("C."), align=PP_ALIGN.CENTER)

    label_w = 0.63
    track_x = x0 + 0.77
    track_w = 3.15
    he_y, other_y, track_h = 2.40, 3.12, 0.58
    add_rect(slide, x0 + 0.10, he_y + 0.12, label_w, 0.34, TEAL_PALE, edge=TEAL_PALE, rounded=True)
    add_rect(slide, x0 + 0.10, other_y + 0.12, label_w, 0.34, ORANGE_PALE, edge=ORANGE_PALE, rounded=True)
    add_text(slide, "HE-Head", x0 + 0.11, he_y + 0.12, label_w - 0.02, 0.34, size=7.6, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Other Head", x0 + 0.11, other_y + 0.12, label_w - 0.02, 0.34, size=7.3, fill=ORANGE_DARK, bold=True, align=PP_ALIGN.CENTER)

    evidence_index = 3
    evidence_center = (evidence_index + 0.5) / len(frame_paths)
    he, other = conceptual_curves(34, evidence_center)
    draw_editable_track(slide, track_x, he_y, track_w, track_h, he, TEAL, TEAL_PALE)
    draw_editable_track(slide, track_x, other_y, track_w, track_h, other, ORANGE, ORANGE_PALE)

    frame_y = 3.88
    gap = 0.025
    frame_w = (track_w - gap * (len(frame_paths) - 1)) / len(frame_paths)
    frame_h = 0.54
    centers = []
    for index, (path, timestamp) in enumerate(zip(frame_paths, times)):
        left = track_x + index * (frame_w + gap)
        centers.append(left + frame_w / 2)
        add_picture_crop(slide, path, left, frame_y, frame_w, frame_h)
        outline = add_rect(
            slide,
            left,
            frame_y,
            frame_w,
            frame_h,
            WHITE,
            edge=TEAL if index == evidence_index else WHITE,
            width=2.0 if index == evidence_index else 0.45,
        )
        outline.fill.transparency = 100
        add_text(slide, str(timestamp), left, frame_y + frame_h + 0.02, frame_w, 0.16, size=4.8, fill=MUTED, align=PP_ALIGN.CENTER)
    evidence_x = centers[evidence_index]
    dashed_vertical(slide, evidence_x, he_y - 0.10, frame_y + frame_h, fill=TEAL)
    add_text(slide, "evidence", evidence_x - 0.31, he_y - 0.23, 0.62, 0.18, size=6.6, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "video history", x0 + 0.09, frame_y + 0.08, 0.65, 0.30, size=6.7, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, track_x, 4.70, track_x + track_w - 0.12, 4.70, fill=MUTED, width=0.9)
    add_arrow(slide, track_x + track_w - 0.18, 4.64, 0.18, 0.12, fill=MUTED)
    add_line(slide, track_x + track_w, 4.43, track_x + track_w, 4.87, fill=BLUE, width=1.7)
    add_text(slide, "tq", track_x + track_w - 0.10, 4.86, 0.20, 0.17, size=6.2, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HE-Head -> C", x0 + 0.25, 5.15, 1.50, 0.26, size=8.0, fill=TEAL_DARK, bold=True)
    add_text(slide, "Other Head -> misses evidence", x0 + 1.83, 5.15, 2.05, 0.26, size=7.8, fill=ORANGE_DARK, bold=True, align=PP_ALIGN.RIGHT)


def draw_panel_b(slide, x0, scores):
    panel_title(slide, "b", "Discovering HE-Heads", x0 + 0.10, 0.82, TEAL)
    add_text(slide, "True-question attention", x0 + 0.18, 1.25, 1.42, 0.24, size=7.2, fill=NAVY, bold=True)
    add_rect(slide, x0 + 0.19, 1.61, 0.48, 0.30, ORANGE_PALE, edge=ORANGE_PALE, rounded=True)
    add_rect(slide, x0 + 0.67, 1.61, 1.34, 0.30, TEAL_PALE, edge=TEAL_PALE)
    add_rect(slide, x0 + 2.01, 1.61, 0.48, 0.30, ORANGE_PALE, edge=ORANGE_PALE)
    add_rect(slide, x0 + 2.49, 1.61, 0.83, 0.30, "E4EEF6", edge="E4EEF6", rounded=True)
    add_text(slide, "start", x0 + 0.19, 1.61, 0.48, 0.30, size=5.8, fill=ORANGE_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "clean internal history", x0 + 0.69, 1.61, 1.30, 0.30, size=6.2, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "boundary", x0 + 2.01, 1.61, 0.48, 0.30, size=5.5, fill=ORANGE_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "current", x0 + 2.49, 1.61, 0.83, 0.30, size=6.0, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, x0 + 0.43, 1.95, x0 + 0.43, 2.18, fill=CORAL, width=1.1)
    add_line(slide, x0 + 2.25, 1.95, x0 + 2.25, 2.18, fill=CORAL, width=1.1)
    add_text(slide, "exclude", x0 + 0.12, 2.14, 0.65, 0.20, size=5.8, fill=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "exclude", x0 + 1.94, 2.14, 0.65, 0.20, size=5.8, fill=CORAL, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, x0 + 3.42, 1.68, 0.42, 0.15, fill=LINE)

    add_rect(slide, x0 + 0.36, 2.54, 3.31, 0.66, WHITE, edge=TEAL, rounded=True, width=0.9)
    add_text(
        slide,
        "Mg = Eq [ Top100Mean ( Agq ( Hinternal ) ) ]",
        x0 + 0.49,
        2.60,
        3.05,
        0.31,
        size=8.3,
        fill=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "boundary-cleaned historical readout", x0 + 0.70, 2.92, 2.63, 0.19, size=6.0, fill=MUTED, align=PP_ALIGN.CENTER)

    matrix = np.full((28, 4), np.nan)
    for row in scores.itertuples():
        matrix[int(row.layer), int(row.kv_head)] = float(row.internal_top100_mean_attention)
    low, high = np.nanquantile(matrix, [0.05, 0.95])
    threshold = np.nanquantile(matrix, 0.90)
    grid_x, grid_y = x0 + 0.57, 3.47
    cell_w, cell_h = 0.34, 0.092
    for layer in range(28):
        for head in range(4):
            value = matrix[layer, head]
            shade = mix("FFF0CF", "C94455", (value - low) / max(high - low, 1e-12))
            cell = add_rect(
                slide,
                grid_x + head * cell_w,
                grid_y + layer * cell_h,
                cell_w,
                cell_h,
                shade,
                edge=WHITE,
                width=0.18,
            )
            if value >= threshold:
                cell.line.color.rgb = color(TEAL_DARK)
                cell.line.width = Pt(0.75)
    for head in range(4):
        add_text(slide, f"K{head}", grid_x + head * cell_w, grid_y - 0.20, cell_w, 0.16, size=5.5, fill=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "L0", grid_x - 0.29, grid_y - 0.02, 0.25, 0.16, size=5.4, fill=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "L27", grid_x - 0.31, grid_y + 27 * cell_h - 0.02, 0.27, 0.16, size=5.4, fill=MUTED, align=PP_ALIGN.RIGHT)

    add_text(slide, "28 layers x 4 physical KV Heads", x0 + 2.18, 3.50, 1.55, 0.34, size=7.2, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, x0 + 2.30, 4.08, 1.35, 0.48, TEAL_PALE, edge=TEAL_PALE, rounded=True)
    add_text(slide, "Historical\nEvidence Heads", x0 + 2.39, 4.10, 1.17, 0.43, size=7.5, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, x0 + 2.28, 5.10, 1.40, 0.78, WHITE, edge=LINE, rounded=True)
    add_text(slide, "Stable on unseen videos", x0 + 2.38, 5.18, 1.20, 0.25, size=6.5, fill=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Spearman rho = 0.938", x0 + 2.34, 5.48, 1.28, 0.26, size=8.0, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)


def cache_rows(slide, x, y, factors, fills, label_prefix="H"):
    max_w = 1.20
    for index, (factor, fill) in enumerate(zip(factors, fills)):
        yy = y + index * 0.48
        add_text(slide, f"{label_prefix}{index + 1}", x, yy, 0.30, 0.25, size=6.4, fill=MUTED, align=PP_ALIGN.RIGHT)
        add_rect(slide, x + 0.36, yy, max_w * factor, 0.25, fill, edge=fill, rounded=True)
        token_count = max(2, int(round(7 * factor)))
        for token in range(token_count):
            tx = x + 0.41 + token * max(max_w * factor - 0.12, 0.05) / max(token_count - 1, 1)
            add_circle(slide, tx, yy + 0.08, 0.07, WHITE, edge=WHITE, width=0.2)


def draw_panel_c(slide, x0):
    panel_title(slide, "c", "Evidence-aware KV allocation", x0 + 0.10, 0.82, ORANGE)
    add_text(slide, "Uniform HERMES", x0 + 0.26, 1.35, 1.40, 0.28, size=8.5, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "MemHeadKV", x0 + 2.43, 1.35, 1.35, 0.28, size=8.5, fill=ORANGE_DARK, bold=True, align=PP_ALIGN.CENTER)
    cache_rows(slide, x0 + 0.16, 1.76, [0.72] * 6, [BLUE] * 6)
    hetero = [1.00, 0.42, 0.88, 0.50, 0.94, 0.58]
    # Sum equals six uniform rows: 4.32.
    cache_rows(slide, x0 + 2.28, 1.76, hetero, [TEAL, ORANGE, TEAL, ORANGE, TEAL, ORANGE])
    add_arrow(slide, x0 + 1.82, 2.82, 0.38, 0.18, fill=LINE)
    add_text(slide, "HE profile", x0 + 1.72, 2.47, 0.58, 0.22, size=6.2, fill=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "same total B", x0 + 1.68, 3.13, 0.66, 0.22, size=6.0, fill=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)

    add_rect(slide, x0 + 0.34, 4.72, 3.38, 0.65, WHITE, edge=LINE, rounded=True)
    add_text(slide, "Offline profile", x0 + 0.43, 4.87, 0.82, 0.24, size=6.6, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, x0 + 1.28, 4.92, 0.31, 0.12, fill=LINE)
    add_text(slide, "Head-wise budget", x0 + 1.62, 4.87, 1.00, 0.24, size=6.6, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, x0 + 2.66, 4.92, 0.31, 0.12, fill=LINE)
    add_text(slide, "Ragged cache", x0 + 3.00, 4.87, 0.62, 0.24, size=6.6, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)

    badges = [
        ("Training-free", TEAL_PALE, TEAL_DARK, 1.02),
        ("Physical KV Heads", ORANGE_PALE, ORANGE_DARK, 1.30),
        ("Same Budget", "E4EEF6", BLUE, 1.02),
    ]
    bx = x0 + 0.34
    for label, fill, text_fill, width in badges:
        add_rect(slide, bx, 5.67, width, 0.34, fill, edge=fill, rounded=True)
        add_text(slide, label, bx + 0.03, 5.68, width - 0.06, 0.31, size=6.4, fill=text_fill, bold=True, align=PP_ALIGN.CENTER)
        bx += width + 0.10
    add_text(
        slide,
        "Protect the Heads that retrieve historical evidence.",
        x0 + 0.34,
        6.25,
        3.38,
        0.32,
        size=8.3,
        fill=ORANGE_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    frame_paths, times = extract_frames()
    scores = pd.read_csv(SCORES_PATH)

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.2)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(WHITE)

    add_text(slide, "Not All Visual Heads Remember", 0.30, 0.14, 5.20, 0.38, size=19, fill=NAVY, bold=True)
    add_text(
        slide,
        "Discovering Historical Evidence Heads for Streaming Video Understanding",
        5.55,
        0.18,
        7.45,
        0.30,
        size=10.5,
        fill=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    add_line(slide, 0.30, 0.60, 13.03, 0.60, fill=LINE, width=0.8)

    panel_x = [0.25, 4.54, 8.83]
    panel_w = [4.05, 4.05, 4.25]
    for x, width in zip(panel_x, panel_w):
        add_rect(slide, x, 0.70, width, 6.18, PANEL, edge=LINE, rounded=True, width=0.65)
    add_line(slide, 4.42, 0.88, 4.42, 6.70, fill=LINE, width=0.7)
    add_line(slide, 8.71, 0.88, 8.71, 6.70, fill=LINE, width=0.7)

    draw_panel_a(slide, panel_x[0], frame_paths, times)
    draw_panel_b(slide, panel_x[1], scores)
    draw_panel_c(slide, panel_x[2])

    presentation.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
