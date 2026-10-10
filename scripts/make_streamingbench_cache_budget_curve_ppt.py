"""Create an editable 4K/6K StreamingBench cache-budget comparison curve."""

from collections import Counter
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULT_ROOT = ROOT / "results/qwen2.5_vl_7b/streamingbench"
OUT_DIR = ROOT / "results/cvpr/figure1"
PPT_PATH = OUT_DIR / "streamingbench_cache_budget_curve.pptx"
PREVIEW_PATH = OUT_DIR / "streamingbench_cache_budget_curve_preview.png"

FILES = {
    "baseline_4000": RESULT_ROOT / "fps0.5-kv4000-hermes/results.csv",
    "baseline_6000": RESULT_ROOT / "fps0.5-kv6000-hermes/results.csv",
    "ours_4000": RESULT_ROOT / "fps0.5-kv4000-hermes-kvheadbudget-v8_blend3_scores-sparsemm-union1-r0.2-w32-raggedprefill-deferlast-strict/results.csv",
    "ours_6000": RESULT_ROOT / "fps0.5-kv6000-hermes-kvheadbudget-v8_blend3_scores-sparsemm-union1-r0.2-w32-raggedprefill-deferlast-strict/results.csv",
}

NAVY, INK, MUTED = "19334C", "2B3943", "71808A"
BLUE, TEAL, TEAL_DARK = "4A78B4", "14918A", "086B65"
PANEL, LINE, WHITE = "F3F8FA", "BDD0D7", "FFFFFF"
TEAL_PALE, BLUE_PALE = "DEF3F0", "E4EEF6"


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=10, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(slide, x, y, w, h, fill=WHITE, edge=LINE, rounded=False, width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(edge)
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=0.8):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def occurrence_keys(frame):
    key = ["video_id", "question", "correct_choice"]
    counts = Counter()
    result = []
    for row in frame[key].itertuples(index=False, name=None):
        occurrence = counts[row]
        counts[row] += 1
        result.append((*row, occurrence))
    return result


def aligned_mean(reference, candidate):
    key = ["video_id", "question", "correct_choice"]
    ref = reference.copy()
    cand = candidate.copy()
    ref["occurrence"] = ref.groupby(key).cumcount()
    cand["occurrence"] = cand.groupby(key).cumcount()
    merged = ref[key + ["occurrence"]].merge(
        cand[key + ["occurrence", "qa_acc"]],
        on=key + ["occurrence"],
        how="left",
        validate="one_to_one",
    )
    if merged.qa_acc.isna().any():
        raise ValueError(f"Candidate misses {int(merged.qa_acc.isna().sum())} reference questions")
    return float(merged.qa_acc.mean())


def load_scores():
    data = {name: pd.read_csv(path) for name, path in FILES.items()}
    reference = data["baseline_4000"]
    if len(reference) != 1255:
        raise ValueError(f"Expected 1255-question common subset, got {len(reference)}")
    if Counter(occurrence_keys(reference)) != Counter(occurrence_keys(data["ours_4000"])):
        raise ValueError("4K baseline and ours do not use the same question multiset")
    if Counter(occurrence_keys(reference)) != Counter(occurrence_keys(data["ours_6000"])):
        raise ValueError("6K ours and 4K baseline do not use the same question multiset")
    return {
        "HERMES": [float(reference.qa_acc.mean()), aligned_mean(reference, data["baseline_6000"])],
        "Ours": [float(data["ours_4000"].qa_acc.mean()), float(data["ours_6000"].qa_acc.mean())],
    }


def add_native_chart(slide, scores):
    data = ChartData()
    data.categories = ["4,000", "6,000"]
    data.add_series("HERMES", scores["HERMES"])
    data.add_series("Ours", scores["Ours"])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.70), Inches(1.25), Inches(8.65), Inches(5.22), data,
    ).chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name = "Arial"
    chart.legend.font.size = Pt(10)
    chart.value_axis.minimum_scale = 72
    chart.value_axis.maximum_scale = 77
    chart.value_axis.major_unit = 1
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.name = "Arial"
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.axis_title.text_frame.paragraphs[0].text = "Average accuracy (%)"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    chart.category_axis.tick_labels.font.name = "Arial"
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.bold = True
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.paragraphs[0].text = "Average KV cache budget"
    chart.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    chart.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)

    colors = [BLUE, TEAL]
    for series, color in zip(chart.series, colors):
        series.format.line.color.rgb = rgb(color)
        series.format.line.width = Pt(2.6)
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = rgb(color)
        series.marker.format.line.color.rgb = rgb(WHITE)
        series.marker.size = 10
        series.has_data_labels = True
        labels = series.data_labels
        labels.position = XL_LABEL_POSITION.ABOVE
        labels.number_format = "0.00"
        labels.font.name = "Arial"
        labels.font.size = Pt(10)
        labels.font.bold = True
        labels.font.color.rgb = rgb(color)
    return chart


def add_summary(slide, scores):
    x, y, w, h = 9.62, 1.25, 3.20, 5.22
    add_rect(slide, x, y, w, h, PANEL, LINE, rounded=True)
    add_text(slide, "Gain over HERMES", x + 0.22, y + 0.20, 2.76, 0.32, size=12, color=NAVY, bold=True)
    add_text(slide, "same 1,255 questions", x + 0.22, y + 0.51, 2.76, 0.22, size=8, color=MUTED)
    add_line(slide, x + 0.22, y + 0.84, x + 2.98, y + 0.84)
    for idx, budget in enumerate([4000, 6000]):
        yy = y + 1.14 + idx * 1.35
        gain = scores["Ours"][idx] - scores["HERMES"][idx]
        add_text(slide, f"{budget:,} budget", x + 0.28, yy, 1.25, 0.30, size=9, color=INK, bold=True)
        add_rect(slide, x + 1.70, yy - 0.02, 1.08, 0.43, TEAL_PALE, TEAL_PALE, rounded=True)
        add_text(slide, f"+{gain:.2f}", x + 1.76, yy, 0.96, 0.38, size=13, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, f'{scores["HERMES"][idx]:.2f}  ->  {scores["Ours"][idx]:.2f}', x + 0.29, yy + 0.49, 2.46, 0.27, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, x + 0.24, y + 4.09, 2.72, 0.70, WHITE, LINE, rounded=True)
    add_text(slide, "0.5 FPS  |  Qwen2.5-VL-7B", x + 0.38, y + 4.20, 2.44, 0.22, size=7.7, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "v8_blend3 + strict cache update", x + 0.38, y + 4.44, 2.44, 0.20, size=7.0, color=MUTED, align=PP_ALIGN.CENTER)


def make_ppt(scores):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.2)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_text(slide, "StreamingBench Accuracy vs. Cache Budget", 0.38, 0.15, 7.65, 0.43, size=20, color=NAVY, bold=True)
    add_text(slide, "Common 1,255-QA subset", 8.85, 0.23, 4.08, 0.27, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_line(slide, 0.38, 0.66, 12.95, 0.66)
    add_native_chart(slide, scores)
    add_summary(slide, scores)
    add_text(slide, "Y-axis is zoomed to 72-77 for readability; full 2,500-question 4K evaluation is pending.", 0.55, 6.79, 12.20, 0.22, size=7.4, color=MUTED, align=PP_ALIGN.CENTER)
    prs.save(PPT_PATH)


def make_preview(scores):
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(13.333, 7.2), dpi=180, facecolor="white")
    ax = fig.add_axes([0.085, 0.19, 0.60, 0.65])
    budgets = [4000, 6000]
    for label, values, color in [("HERMES", scores["HERMES"], "#" + BLUE), ("Ours", scores["Ours"], "#" + TEAL)]:
        ax.plot(budgets, values, marker="o", markersize=9, linewidth=3, color=color, label=label)
        for x, value in zip(budgets, values):
            ax.text(x, value + 0.16, f"{value:.2f}", ha="center", fontsize=11, fontweight="bold", color=color)
    ax.set_xlim(3700, 6300); ax.set_xticks(budgets, ["4,000", "6,000"])
    ax.set_ylim(72, 77); ax.set_yticks([72, 73, 74, 75, 76, 77])
    ax.set_xlabel("Average KV cache budget", fontsize=11, fontweight="bold", color="#" + INK)
    ax.set_ylabel("Average accuracy (%)", fontsize=11, fontweight="bold", color="#" + INK)
    ax.grid(axis="y", color="#dbe5e9", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(loc="lower right", frameon=False, fontsize=10)
    fig.text(0.035, 0.93, "StreamingBench Accuracy vs. Cache Budget", fontsize=20, fontweight="bold", color="#" + NAVY)
    fig.text(0.96, 0.94, "Common 1,255-QA subset", ha="right", fontsize=9.5, color="#" + MUTED)
    panel = fig.add_axes([0.73, 0.19, 0.235, 0.65]); panel.axis("off")
    panel.add_patch(plt.Rectangle((0, 0), 1, 1, facecolor="#" + PANEL, edgecolor="#" + LINE, linewidth=1))
    panel.text(0.08, 0.91, "Gain over HERMES", fontsize=12, fontweight="bold", color="#" + NAVY)
    for idx, budget in enumerate(budgets):
        y = 0.70 - idx * 0.27; gain=scores["Ours"][idx]-scores["HERMES"][idx]
        panel.text(0.10, y, f"{budget:,} budget", fontsize=9, fontweight="bold", color="#" + INK)
        panel.text(0.87, y, f"+{gain:.2f}", ha="right", fontsize=15, fontweight="bold", color="#" + TEAL_DARK)
        panel.text(0.50, y-0.08, f'{scores["HERMES"][idx]:.2f}  ->  {scores["Ours"][idx]:.2f}', ha="center", fontsize=8.5, color="#" + MUTED)
    panel.text(0.50, 0.12, "0.5 FPS | Qwen2.5-VL-7B", ha="center", fontsize=8, fontweight="bold", color="#" + NAVY)
    panel.text(0.50, 0.07, "v8_blend3 + strict cache update", ha="center", fontsize=7, color="#" + MUTED)
    fig.text(0.5, 0.04, "Y-axis zoomed to 72-77; full 2,500-question 4K evaluation is pending.", ha="center", fontsize=7.5, color="#" + MUTED)
    fig.savefig(PREVIEW_PATH, facecolor="white")
    plt.close(fig)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    scores = load_scores()
    make_ppt(scores)
    make_preview(scores)
    print(PPT_PATH)
    print(PREVIEW_PATH)
    print(scores)


if __name__ == "__main__":
    main()
