"""Create editable prevalence and cross-video stability plots for HE-Heads."""

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "results/myobservations/effective_memory_readout_core_top100_n50_o5000/raw_effective_readout.csv"
OUT_DIR = ROOT / "results/cvpr/figure1"
PPT_PATH = OUT_DIR / "he_head_prevalence_stability.pptx"
PREVIEW_PATH = OUT_DIR / "he_head_prevalence_stability_preview.png"
METRIC = "internal_top100_mean_attention"
TOP_K = 10

NAVY = "19334C"
INK = "2B3943"
MUTED = "71808A"
TEAL = "14918A"
TEAL_DARK = "086B65"
TEAL_PALE = "DEF3F0"
ORANGE = "E98247"
ORANGE_PALE = "FBE7D8"
GRAY = "B8C4CA"
GRAY_PALE = "E7ECEF"
PANEL = "F3F8FA"
LINE = "BDD0D7"
WHITE = "FFFFFF"


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=10, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(slide, x, y, w, h, fill=WHITE, edge=LINE, radius=False, width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(edge)
    shape.line.width = Pt(width)
    return shape


def add_circle(slide, cx, cy, diameter, fill, edge=None, width=0.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(cx - diameter / 2),
        Inches(cy - diameter / 2),
        Inches(diameter),
        Inches(diameter),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(edge or fill)
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1, y1, x2, y2, color=MUTED, width=0.8, dash=None):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def rank(values):
    return pd.Series(values).rank(method="average").to_numpy(float)


def load_profiles():
    usecols = ["video_idx", "layer", "kv_head", METRIC]
    raw = pd.read_csv(RAW, usecols=usecols)
    videos = np.array(sorted(raw.video_idx.astype(int).unique()))
    midpoint = len(videos) // 2

    def profile(selected):
        return (
            raw[raw.video_idx.isin(selected)]
            .groupby(["layer", "kv_head"], as_index=False)[METRIC]
            .mean()
        )

    full = profile(videos)
    left = profile(videos[:midpoint]).rename(columns={METRIC: "calibration"})
    right = profile(videos[midpoint:]).rename(columns={METRIC: "evaluation"})
    split = left.merge(right, on=["layer", "kv_head"])
    rho = float(np.corrcoef(rank(split.calibration), rank(split.evaluation))[0, 1])

    full_top = set(map(tuple, full.nlargest(TOP_K, METRIC)[["layer", "kv_head"]].to_numpy(int)))
    left_top = set(map(tuple, split.nlargest(TOP_K, "calibration")[["layer", "kv_head"]].to_numpy(int)))
    right_top = set(map(tuple, split.nlargest(TOP_K, "evaluation")[["layer", "kv_head"]].to_numpy(int)))
    return full, split, full_top, left_top, right_top, rho


def panel_header(slide, label, title, x, y, accent):
    add_circle(slide, x + 0.15, y + 0.15, 0.30, accent)
    add_text(slide, label, x, y, 0.30, 0.30, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.39, y - 0.01, 5.55, 0.33, size=12, color=NAVY, bold=True)


def add_pie(slide, x0, full_top):
    chart_data = ChartData()
    chart_data.categories = ["HE-Heads", "Other Heads"]
    chart_data.add_series("Physical KV Heads", [len(full_top), 112 - len(full_top)])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(x0 + 0.38), Inches(1.55), Inches(3.45), Inches(3.45), chart_data
    ).chart
    chart.has_legend = False
    chart.has_title = False
    chart.chart_style = 10
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_DATA_LABEL_POSITION.BEST_FIT
    labels.show_percentage = True
    labels.show_category_name = False
    labels.show_value = False
    labels.font.name = "Arial"
    labels.font.size = Pt(11)
    labels.font.bold = True
    labels.font.color.rgb = rgb(WHITE)
    series = plot.series[0]
    series.points[0].format.fill.solid()
    series.points[0].format.fill.fore_color.rgb = rgb(TEAL)
    series.points[1].format.fill.solid()
    series.points[1].format.fill.fore_color.rgb = rgb(GRAY)
    for point in series.points:
        point.format.line.color.rgb = rgb(WHITE)
        point.format.line.width = Pt(1.2)

    add_text(slide, "10", x0 + 4.05, 1.90, 1.05, 0.62, size=27, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HE-Heads", x0 + 3.92, 2.45, 1.30, 0.30, size=10, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "8.9% of 112", x0 + 3.92, 2.77, 1.30, 0.27, size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, x0 + 4.10, 3.36, 0.25, 0.16, TEAL, TEAL, radius=True)
    add_text(slide, "HE-Heads", x0 + 4.48, 3.29, 1.05, 0.30, size=9, color=INK)
    add_rect(slide, x0 + 4.10, 3.82, 0.25, 0.16, GRAY, GRAY, radius=True)
    add_text(slide, "Other Heads", x0 + 4.48, 3.75, 1.20, 0.30, size=9, color=INK)
    add_rect(slide, x0 + 0.65, 5.40, 4.95, 0.67, TEAL_PALE, TEAL_PALE, radius=True)
    add_text(
        slide,
        "Top-10 boundary-cleaned historical readout",
        x0 + 0.78,
        5.51,
        4.69,
        0.26,
        size=9.5,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_scatter(slide, x0, split, left_top, right_top, rho):
    plot_x, plot_y, plot_w, plot_h = x0 + 0.78, 1.53, 4.90, 4.68
    plot_data = split.copy()
    plot_data["calibration_rank"] = plot_data.calibration.rank(method="average", pct=True)
    plot_data["evaluation_rank"] = plot_data.evaluation.rank(method="average", pct=True)
    axis_min, axis_max = 0.0, 1.0

    def px(value):
        return plot_x + (value - axis_min) / (axis_max - axis_min) * plot_w

    def py(value):
        return plot_y + plot_h - (value - axis_min) / (axis_max - axis_min) * plot_h

    add_rect(slide, plot_x, plot_y, plot_w, plot_h, WHITE, "D8E4E8", width=0.8)
    ticks = np.linspace(axis_min, axis_max, 5)
    for tick in ticks:
        tx, ty = px(tick), py(tick)
        add_line(slide, tx, plot_y, tx, plot_y + plot_h, color="E5EBEE", width=0.45)
        add_line(slide, plot_x, ty, plot_x + plot_w, ty, color="E5EBEE", width=0.45)
        add_text(slide, f"{tick:.2f}", tx - 0.22, plot_y + plot_h + 0.04, 0.44, 0.22, size=6.7, color=MUTED, align=PP_ALIGN.CENTER)
        add_text(slide, f"{tick:.2f}", plot_x - 0.48, ty - 0.11, 0.39, 0.22, size=6.7, color=MUTED, align=PP_ALIGN.RIGHT)

    add_line(slide, px(axis_min), py(axis_min), px(axis_max), py(axis_max), color="93A8B0", width=1.0)
    for row in plot_data.itertuples():
        key = (int(row.layer), int(row.kv_head))
        shared = key in left_top and key in right_top
        split_only = key in left_top or key in right_top
        fill = TEAL if shared else ORANGE if split_only else GRAY
        edge = TEAL_DARK if shared else ORANGE if split_only else WHITE
        size = 0.115 if shared else 0.095 if split_only else 0.065
        add_circle(slide, px(row.calibration_rank), py(row.evaluation_rank), size, fill, edge, 0.55)

    add_text(slide, "Calibration split HE-score percentile", plot_x + 1.02, plot_y + plot_h + 0.33, 2.90, 0.28, size=8.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    y_label = add_text(slide, "Evaluation split HE-score percentile", plot_x - 0.78, plot_y + 1.14, 0.35, 2.40, size=8.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    y_label.rotation = 270

    add_rect(slide, plot_x + 0.18, plot_y + 0.18, 1.78, 0.65, WHITE, LINE, radius=True)
    add_text(slide, f"Spearman rho = {rho:.3f}", plot_x + 0.28, plot_y + 0.23, 1.58, 0.24, size=9.5, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Top-10 overlap = 8 / 10", plot_x + 0.28, plot_y + 0.49, 1.58, 0.20, size=7.6, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

    legend_y = 6.48
    add_circle(slide, x0 + 1.15, legend_y, 0.10, TEAL, TEAL_DARK)
    add_text(slide, "shared Top-10", x0 + 1.27, legend_y - 0.13, 1.24, 0.26, size=7.5, color=INK)
    add_circle(slide, x0 + 2.72, legend_y, 0.10, ORANGE, ORANGE)
    add_text(slide, "split-specific Top-10", x0 + 2.84, legend_y - 0.13, 1.62, 0.26, size=7.5, color=INK)
    add_circle(slide, x0 + 4.76, legend_y, 0.08, GRAY, WHITE)
    add_text(slide, "other heads", x0 + 4.87, legend_y - 0.13, 1.00, 0.26, size=7.5, color=INK)


def make_ppt(full_top, split, left_top, right_top, rho):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.2)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_text(slide, "Historical-Evidence Heads Are Sparse and Stable", 0.35, 0.15, 8.50, 0.42, size=20, color=NAVY, bold=True)
    add_text(slide, "Qwen2.5-VL-7B  |  StreamingBench  |  50 videos", 8.72, 0.22, 4.25, 0.28, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_line(slide, 0.35, 0.66, 12.98, 0.66, color=LINE, width=0.8)
    add_rect(slide, 0.28, 0.80, 6.18, 6.08, PANEL, LINE, radius=True)
    add_rect(slide, 6.65, 0.80, 6.40, 6.08, PANEL, LINE, radius=True)
    panel_header(slide, "b", "HE-Head prevalence", 0.47, 0.98, TEAL)
    panel_header(slide, "c", "Cross-video stability", 6.84, 0.98, TEAL)
    add_pie(slide, 0.32, full_top)
    add_scatter(slide, 6.61, split, left_top, right_top, rho)
    prs.save(PPT_PATH)


def make_preview(full_top, split, left_top, right_top, rho):
    scale = 180
    width, height = 2400, 1296
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_dir = Path(__import__("matplotlib").get_data_path()) / "fonts" / "ttf"

    def f(size, bold=False):
        return ImageFont.truetype(str(font_dir / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")), size)

    draw.text((55, 38), "Historical-Evidence Heads Are Sparse and Stable", font=f(39, True), fill="#" + NAVY)
    draw.rounded_rectangle((50, 145, 1145, 1225), 24, fill="#" + PANEL, outline="#" + LINE, width=2)
    draw.rounded_rectangle((1195, 145, 2350, 1225), 24, fill="#" + PANEL, outline="#" + LINE, width=2)
    draw.text((105, 195), "(b) HE-Head prevalence", font=f(25, True), fill="#" + NAVY)
    draw.text((1250, 195), "(c) Cross-video stability", font=f(25, True), fill="#" + NAVY)

    center, radius = (455, 615), 285
    start = -90
    he_angle = 360 * len(full_top) / 112
    draw.pieslice((center[0]-radius, center[1]-radius, center[0]+radius, center[1]+radius), start, start + he_angle, fill="#" + TEAL, outline="white", width=3)
    draw.pieslice((center[0]-radius, center[1]-radius, center[0]+radius, center[1]+radius), start + he_angle, start + 360, fill="#" + GRAY, outline="white", width=3)
    draw.text((820, 460), "10", font=f(70, True), fill="#" + TEAL_DARK, anchor="mm")
    draw.text((820, 530), "HE-Heads", font=f(25, True), fill="#" + TEAL_DARK, anchor="mm")
    draw.text((820, 575), "8.9% of 112", font=f(22), fill="#" + MUTED, anchor="mm")
    draw.rounded_rectangle((180, 1040, 1015, 1135), 18, fill="#" + TEAL_PALE)
    draw.text((598, 1088), "Top-10 boundary-cleaned historical readout", font=f(22, True), fill="#" + TEAL_DARK, anchor="mm")

    px0, py0, pw, ph = 1320, 330, 820, 720
    plot_data = split.copy()
    plot_data["calibration_rank"] = plot_data.calibration.rank(method="average", pct=True)
    plot_data["evaluation_rank"] = plot_data.evaluation.rank(method="average", pct=True)
    for tick in np.linspace(0, 1, 5):
        x = px0 + tick*pw; y = py0+ph-tick*ph
        draw.line((x,py0,x,py0+ph),fill="#e5ebee",width=2); draw.line((px0,y,px0+pw,y),fill="#e5ebee",width=2)
    draw.rectangle((px0,py0,px0+pw,py0+ph),outline="#"+LINE,width=2)
    draw.line((px0,py0+ph,px0+pw,py0),fill="#93a8b0",width=3)
    for row in plot_data.itertuples():
        key=(int(row.layer),int(row.kv_head)); shared=key in left_top and key in right_top; special=key in left_top or key in right_top
        color="#"+(TEAL if shared else ORANGE if special else GRAY); radius=8 if shared else 6 if special else 4
        x=px0+row.calibration_rank*pw; y=py0+ph-row.evaluation_rank*ph
        draw.ellipse((x-radius,y-radius,x+radius,y+radius),fill=color,outline="white")
    draw.rounded_rectangle((px0+25,py0+25,px0+335,py0+130),16,fill="white",outline="#"+LINE,width=2)
    draw.text((px0+180,py0+58),f"Spearman rho = {rho:.3f}",font=f(21,True),fill="#"+TEAL_DARK,anchor="mm")
    draw.text((px0+180,py0+100),"Top-10 overlap = 8 / 10",font=f(17,True),fill="#"+MUTED,anchor="mm")
    draw.text((px0+pw/2,py0+ph+60),"Calibration split HE-score percentile",font=f(20,True),fill="#"+INK,anchor="mm")
    image.save(PREVIEW_PATH, dpi=(300,300), optimize=True)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    full, split, full_top, left_top, right_top, rho = load_profiles()
    assert len(full) == 112 and len(full_top) == TOP_K
    assert len(left_top & right_top) == 8
    make_ppt(full_top, split, left_top, right_top, rho)
    make_preview(full_top, split, left_top, right_top, rho)
    print(PPT_PATH)
    print(PREVIEW_PATH)
    print({"he_heads": len(full_top), "fraction": len(full_top) / len(full), "rho": rho, "top10_overlap": len(left_top & right_top)})


if __name__ == "__main__":
    main()
